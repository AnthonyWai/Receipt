"""Extract receipt details from Processing/ into a daily Excel workbook."""

from __future__ import annotations

import base64
import json
import os
import re
import shutil
import sys
from io import BytesIO
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener
import requests

register_heif_opener()


ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "Processing"
OUTPUT_DIR = ROOT / "Processed"
API_URL = "https://api.deepseek.com/chat/completions"
MODEL = "deepseek-chat"
VISION_MODEL = "deepseek-flash"
MAX_VISION_IMAGE_EDGE = 1600
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".heic", ".tif", ".tiff"}
RECEIPT_FIELDS = ("shop_name", "category", "date", "total_price", "currency")


def read_receipt_text(image_path: Path) -> str:
    """Read printed receipt text using the platform's OCR engine."""
    if sys.platform == "darwin":
        try:
            import Vision
            from Foundation import NSURL
        except ImportError as exc:
            raise RuntimeError(
                "macOS Vision OCR dependencies are missing; install requirements.txt."
            ) from exc

        request = Vision.VNRecognizeTextRequest.alloc().init()
        request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
        request.setRecognitionLanguages_(["zh-Hant", "en-US"])
        request.setUsesLanguageCorrection_(True)

        image_url = NSURL.fileURLWithPath_(str(image_path))
        handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(image_url, {})
        succeeded, error = handler.performRequests_error_([request], None)
        if not succeeded:
            raise RuntimeError(f"Vision OCR failed: {error}")

        lines = []
        for observation in request.results() or []:
            candidates = observation.topCandidates_(1)
            if candidates:
                lines.append(candidates[0].string())
        receipt_text = "\n".join(lines)
    else:
        import pytesseract

        try:
            with Image.open(image_path) as image:
                image = ImageOps.exif_transpose(image).convert("RGB")
                receipt_text = pytesseract.image_to_string(
                    image, lang="chi_tra+eng"
                )
        except pytesseract.TesseractNotFoundError as exc:
            raise RuntimeError(
                "Tesseract OCR is required; install Tesseract and add it to PATH."
            ) from exc
        except pytesseract.TesseractError as exc:
            raise RuntimeError(f"Tesseract OCR failed: {exc}") from exc

    lines = [line.strip() for line in receipt_text.splitlines() if line.strip()]
    if not lines:
        raise RuntimeError("No text was recognized in the receipt image.")
    return "\n".join(lines)


def is_unknown(value: Any) -> bool:
    return str(value or "").strip().casefold() in {
        "",
        "unknown",
        "n/a",
        "na",
        "none",
        "null",
    }


def parse_response_fields(response: requests.Response) -> dict[str, Any]:
    response.raise_for_status()
    content = response.json()["choices"][0]["message"]["content"]
    fields = json.loads(content)
    if not isinstance(fields, dict):
        raise ValueError("DeepSeek returned a JSON value that was not an object.")
    return fields


def extract_receipt_fields(receipt_text: str, api_key: str) -> dict[str, Any]:
    """Ask DeepSeek to normalize OCR text into receipt fields."""
    prompt = (
        "Extract the receipt's shop name, purchase category, transaction date, "
        "total paid, and currency from the OCR text. Use an ISO date (YYYY-MM-DD). "
        "Choose a concise category such as groceries, dining, pharmacy, clothing, "
        "transport, utilities, electronics, household, entertainment, or other. "
        "For Taste (including stylized OCR readings such as ToSTe), normalize "
        "shop_name to 'Taste'. "
        "Do not mistake subtotals, discounts, or change for the final amount paid. "
        'Return only a JSON object with keys "shop_name", "category", "date", '
        '"total_price", and "currency". Use empty strings for unknown values.\n\n'
        f"Receipt OCR text:\n{receipt_text}"
    )
    response = requests.post(
        API_URL,
        headers={"Authorization": f"Bearer {api_key}"},
        json={
            "model": MODEL,
            "messages": [
                {"role": "system", "content": "You extract fields from receipts."},
                {"role": "user", "content": prompt},
            ],
            "response_format": {"type": "json_object"},
        },
        timeout=60,
    )
    fields = parse_response_fields(response)
    if "taste" in receipt_text.casefold() or "toste" in receipt_text.casefold():
        fields["shop_name"] = "Taste"
    return fields


def resized_image_data_url(image_path: Path) -> str:
    """Convert and resize a receipt to a compact JPEG data URL."""
    try:
        with Image.open(image_path) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            image.thumbnail((MAX_VISION_IMAGE_EDGE, MAX_VISION_IMAGE_EDGE))
            buffer = BytesIO()
            image.save(buffer, format="JPEG", quality=85)
            image_data = base64.b64encode(buffer.getvalue()).decode("ascii")
    except OSError as exc:
        raise RuntimeError(f"Could not resize receipt image for DeepSeek: {exc}") from exc
    return f"data:image/jpeg;base64,{image_data}"


def extract_receipt_fields_from_image(
    image_path: Path, receipt_text: str, api_key: str
) -> dict[str, Any]:
    """Use DeepSeek Vision when OCR-text extraction misses key receipt fields."""
    prompt = (
        "Read this receipt image and extract shop name, purchase category, "
        "transaction date, total paid, and currency. Use the image as the source "
        "of truth; the OCR text is only a hint. Use an ISO date (YYYY-MM-DD). "
        "Do not confuse subtotal, discounts, coupons, or change with total paid. "
        'Return only a JSON object with keys "shop_name", "category", "date", '
        '"total_price", and "currency". Use empty strings when a value is unclear.\n\n'
        f"Local OCR text:\n{receipt_text}"
    )
    response = requests.post(
        API_URL,
        headers={"Authorization": f"Bearer {api_key}"},
        json={
            "model": VISION_MODEL,
            "messages": [
                {"role": "system", "content": "You extract fields from receipts."},
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": resized_image_data_url(image_path),
                                "detail": "original",
                            },
                        },
                    ],
                },
            ],
            "response_format": {"type": "json_object"},
        },
        timeout=90,
    )
    return parse_response_fields(response)


def extract_with_fallback(
    image_path: Path, receipt_text: str, api_key: str
) -> dict[str, Any]:
    fields = extract_receipt_fields(receipt_text, api_key)
    if is_unknown(fields.get("shop_name")) or is_unknown(fields.get("total_price")):
        vision_fields = extract_receipt_fields_from_image(image_path, receipt_text, api_key)
        for field in RECEIPT_FIELDS:
            if is_unknown(fields.get(field)) and not is_unknown(vision_fields.get(field)):
                fields[field] = vision_fields[field]
        fields["_models_used"] = f"{MODEL}, {VISION_MODEL}"
    else:
        fields["_models_used"] = MODEL
    if "taste" in receipt_text.casefold() or "toste" in receipt_text.casefold():
        fields["shop_name"] = "Taste"
    return fields


def receipt_date(fields: dict[str, Any], image_path: Path) -> date:
    value = str(fields.get("date", "")).strip()
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return datetime.fromtimestamp(image_path.stat().st_mtime).date()


def next_processed_image_path(image_path: Path) -> Path:
    processed_images = [
        path
        for path in OUTPUT_DIR.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    ]
    numbered_ids = []
    unnumbered_count = 0
    for processed_image in processed_images:
        match = re.search(r"_(\d+)$", processed_image.stem)
        if match:
            numbered_ids.append(int(match.group(1)))
        else:
            unnumbered_count += 1

    sequence_id = max(max(numbered_ids, default=0), unnumbered_count) + 1
    destination = OUTPUT_DIR / f"{image_path.stem}_{sequence_id:04d}{image_path.suffix}"
    while destination.exists():
        sequence_id += 1
        destination = OUTPUT_DIR / f"{image_path.stem}_{sequence_id:04d}{image_path.suffix}"
    return destination


def save_receipt(
    fields: dict[str, Any], image_path: Path, processed_image_name: str
) -> Path:
    transaction_date = receipt_date(fields, image_path)
    total = str(fields.get("total_price", "")).strip() or "Unknown"
    currency = str(fields.get("currency", "")).strip()
    price = f"{currency} {total}".strip()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / f"{date.today().isoformat()}.xlsx"

    if output_path.exists():
        workbook = load_workbook(output_path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Receipts"
        worksheet.append(
            ["Shop Name", "Category", "Receipt Date", "Price", "AI Model", "Source"]
        )
        worksheet.freeze_panes = "A2"

    worksheet.append(
        [
            str(fields.get("shop_name", "")).strip() or "Unknown",
            str(fields.get("category", "")).strip() or "Other",
            transaction_date,
            price,
            str(fields.get("_models_used", MODEL)),
            processed_image_name,
        ]
    )
    worksheet.cell(row=worksheet.max_row, column=3).number_format = "yyyy-mm-dd"
    workbook.save(output_path)
    workbook.close()
    return output_path


def process_receipt(image_path: Path, api_key: str) -> tuple[Path, Path]:
    receipt_text = read_receipt_text(image_path)
    fields = extract_with_fallback(image_path, receipt_text, api_key)
    processed_image_path = next_processed_image_path(image_path)
    output_path = save_receipt(fields, image_path, processed_image_path.name)
    shutil.move(str(image_path), str(processed_image_path))
    return output_path, processed_image_path


def main() -> int:
    api_key = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if not api_key:
        raise SystemExit("Set the DEEPSEEK_API_KEY environment variable first.")
    if not INPUT_DIR.is_dir():
        raise SystemExit(f"Input folder not found: {INPUT_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    image_paths = sorted(
        path for path in INPUT_DIR.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    )
    if not image_paths:
        print(f"No receipt images found in {INPUT_DIR}.")
        return 0

    failures = 0
    for image_path in image_paths:
        try:
            output_path, processed_image_path = process_receipt(image_path, api_key)
            print(
                f"Processed {image_path.name} -> {processed_image_path.name} "
                f"({output_path.name})"
            )
        except (requests.RequestException, RuntimeError, ValueError, KeyError) as exc:
            failures += 1
            print(f"Failed {image_path.name}: {exc}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())